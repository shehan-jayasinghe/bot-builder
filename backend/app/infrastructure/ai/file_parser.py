import logging
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.domain.models.source_document import SourceDocument

logger = logging.getLogger(__name__)


def parse_file_bytes(*, data: bytes, file_name: str) -> list[SourceDocument]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(data=data, file_name=file_name)
    if suffix == ".docx":
        return _parse_docx(data=data, file_name=file_name)
    if suffix in {".txt", ".md", ".csv"}:
        return [_parse_text(data=data, file_name=file_name)]
    raise ValueError(f"Unsupported file type: {suffix or file_name}")


def _parse_pdf(*, data: bytes, file_name: str) -> list[SourceDocument]:
    reader = PdfReader(BytesIO(data))
    documents: list[SourceDocument] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        documents.append(
            SourceDocument(
                text=text,
                metadata={"source": file_name, "page": page_number},
            )
        )
    if not documents:
        raise ValueError(f"No extractable text found in {file_name}")
    return documents


def _parse_docx(*, data: bytes, file_name: str) -> list[SourceDocument]:
    document = Document(BytesIO(data))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    text = "\n\n".join(paragraphs).strip()
    if not text:
        raise ValueError(f"No extractable text found in {file_name}")
    return [SourceDocument(text=text, metadata={"source": file_name})]


def _parse_text(*, data: bytes, file_name: str) -> SourceDocument:
    text = data.decode("utf-8", errors="replace").strip()
    if not text:
        raise ValueError(f"No extractable text found in {file_name}")
    return SourceDocument(text=text, metadata={"source": file_name})
